from io import BytesIO
from zipfile import ZipFile
import fitz
import pytest
from docx import Document
from conftest import login
from app.documents.representations import represent
from app.core.exceptions import DomainError


def test_document_read_only_pages_and_word_keep_file_permissions(admin):
    base='/api/documents/test-contract/file'
    assert admin.get(base+'?representation=info').json()=={'pages':1}
    page=admin.get(base+'?representation=png&page=1')
    assert page.status_code==200 and page.content.startswith(b'\x89PNG')
    assert page.headers['cache-control']=='no-store'
    assert admin.get(base+'?representation=png&page=2').status_code==404
    exported=admin.get(base+'?representation=docx')
    document=Document(BytesIO(exported.content))
    assert any('Contract mentenanta' in p.text for p in document.paragraphs)
    login(admin,'technician')
    assert admin.get('/api/documents/test-hr/file?representation=info').status_code==404
    assert admin.get('/api/documents/test-hr/file?representation=png').status_code==404
    assert admin.get('/api/documents/test-hr/file?representation=docx').status_code==404


def test_word_export_uses_only_redacted_copy_and_obeys_revocation(admin):
    preview=admin.post('/api/documents/test-procedure/share-preview',json={'redact_terms':['SecretAccount123']})
    assert preview.status_code==200,preview.text
    created=admin.post('/api/documents/test-procedure/shares',json={'preview_id':preview.json()['preview_id'],
        'reviewed':True,'expires_at':'2026-12-31','recipient_airport_id':'sibiu'})
    assert created.status_code==201,created.text
    url='/api/shares/'+created.json()['id']+'/file'
    response=admin.get(url+'?representation=docx')
    assert response.status_code==200,response.text
    with ZipFile(BytesIO(response.content)) as archive:
        assert all(b'SecretAccount123' not in archive.read(name) for name in archive.namelist())
    assert 'Procedura documentelor' in '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    admin.post('/api/shares/'+created.json()['id']+'/revoke')
    for representation in ('info','png','docx'):
        assert admin.get(url+'?representation='+representation).status_code==403


def test_received_word_and_pages_fetch_only_authorized_partner_copy(admin,monkeypatch,tmp_path,pdf_bytes):
    import httpx
    from app.exchange import service
    from app.config import settings
    token=tmp_path/'token';token.write_text('test-exchange-token')
    monkeypatch.setattr(settings,'exchange_token_file',token)
    monkeypatch.setattr(settings,'exchange_peers',{'sibiu':'http://sibiu.test'})
    calls=[]
    def source(url,**kwargs):
        calls.append(url)
        return httpx.Response(200,content=pdf_bytes,headers={'content-type':'application/pdf'})
    monkeypatch.setattr(service.httpx,'get',source)
    for representation in ('info','png','docx'):
        result=admin.get('/api/received/peer-copy/file?airport_id=sibiu&representation='+representation)
        assert result.status_code==200,result.text
        assert calls[-1]=='http://sibiu.test/outgoing/shares/peer-copy/file'
    monkeypatch.setattr(service.httpx,'get',lambda *a,**k:httpx.Response(403))
    assert admin.get('/api/received/peer-copy/file?airport_id=sibiu&representation=docx').status_code==403


def test_editable_template_contains_real_word_tables():
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    output=BytesIO()
    table=Table([['Piesa','Cantitate','Cost'],['Cauciuc','2','150']])
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),1,colors.black)]))
    SimpleDocTemplate(output).build([Paragraph('Model mentenanta',getSampleStyleSheet()['Heading1']),table])
    response=represent(output.getvalue(),'docx')
    word=Document(BytesIO(response.body))
    assert len(word.tables)==1
    assert word.tables[0].cell(1,0).text=='Cauciuc'
    assert word.tables[0].cell(1,2).text=='150'


def test_empty_scan_does_not_offer_a_fake_editable_document():
    with fitz.open() as pdf:
        pdf.new_page()
        with pytest.raises(DomainError,match='text editabil'):
            represent(pdf.tobytes(),'docx')
