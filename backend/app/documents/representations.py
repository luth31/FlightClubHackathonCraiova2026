"""Read-only previews and editable exports of an already authorized PDF only."""
from io import BytesIO
import re
from typing import Literal
import fitz
from docx import Document
from docx.shared import Pt, Cm
from fastapi.responses import Response, JSONResponse
from app.core.exceptions import DomainError

Representation = Literal['pdf', 'info', 'png', 'docx']
HEADERS = {'Cache-Control': 'no-store', 'X-Content-Type-Options': 'nosniff'}
DOCX_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def clean_text(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value or '')


def editable_copy(pdf):
    if len(pdf) > 100:
        raise DomainError('Exportul Word acceptă maximum 100 de pagini. Descarcă PDF-ul.', 422)
    document = Document()
    document.core_properties.title = 'Copie pentru adaptare locală'
    document.core_properties.author = 'FlyLeaf'
    document.core_properties.comments = 'Derivat din PDF-ul accesibil. Verifică aranjarea și conținutul înainte de utilizare.'
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    section = document.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.footer.paragraphs[0].text = 'Copie pentru adaptare locală · verifică forma finală înainte de aprobare.'
    has_text = False
    for index, page in enumerate(pdf):
        if index: document.add_page_break()
        # Use only this PDF's current text, never the original, OCR cache or metadata.
        tables = page.find_tables().tables
        regions = [fitz.Rect(table.bbox) for table in tables]
        items = [(table.bbox[1], table.bbox[0], 'table', table.extract()) for table in tables]
        for block in page.get_text('dict')['blocks']:
            if block['type'] != 0: continue
            for line in block['lines']:
                bounds = fitz.Rect(line['bbox'])
                center = (bounds.tl + bounds.br) / 2
                if any(center in region for region in regions): continue
                items.append((bounds.y0, bounds.x0, 'line', line['spans']))
        for _, _, kind, value in sorted(items, key=lambda item: item[:2]):
            if kind == 'table':
                columns = max((len(row) for row in value), default=0)
                if not columns: continue
                table = document.add_table(rows=len(value), cols=columns)
                table.style = 'Table Grid'
                for r, row in enumerate(value):
                    for c, text in enumerate(row):
                        text = clean_text(text)
                        table.cell(r, c).text = text
                        has_text = has_text or bool(text.strip())
            else:
                paragraph = document.add_paragraph()
                for span in value:
                    text = clean_text(span['text'])
                    has_text = has_text or bool(text.strip())
                    run = paragraph.add_run(text)
                    run.bold = bool(span['flags'] & 16)
                    run.italic = bool(span['flags'] & 2)
                    run.font.size = Pt(min(24, max(9, span['size'])))
        if not page.get_text().strip():
            document.add_paragraph('[Pagină fără text editabil. Consultă PDF-ul pentru conținutul vizual.]')
    if not has_text:
        raise DomainError('Acest PDF nu conține text editabil. Descarcă PDF-ul pentru prelucrare locală.', 422)
    output = BytesIO(); document.save(output)
    return output.getvalue()


def represent(content, representation: Representation, page_number=1):
    try:
        with fitz.open(stream=content, filetype='pdf') as pdf:
            if not len(pdf) or pdf.needs_pass: raise ValueError('Unreadable PDF')
            if representation == 'info':
                return JSONResponse({'pages': len(pdf)}, headers=HEADERS)
            if representation == 'png':
                if page_number < 1 or page_number > len(pdf):
                    raise DomainError('Pagina nu există.', 404)
                page = pdf[page_number-1]
                if min(page.rect.width, page.rect.height) <= 0: raise ValueError('Invalid page')
                scale = min(1.7, 1800/max(page.rect.width, page.rect.height))
                image = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).tobytes('png')
                return Response(image, media_type='image/png', headers=HEADERS)
            if representation == 'docx':
                return Response(editable_copy(pdf), media_type=DOCX_TYPE,
                    headers={**HEADERS, 'Content-Disposition': 'attachment; filename="Copie-pentru-adaptare.docx"'})
    except DomainError: raise
    except (ValueError, RuntimeError):
        raise DomainError('Documentul nu poate fi previzualizat sau exportat.', 422) from None
    raise DomainError('Format necunoscut.', 422)
